"""Build CV + covering letter for Manchester SAE-032399 (Research Associate in ssNMR).

Usage: python3 build_manchester_ssnmr.py   (run from any directory)
Outputs .docx, .pdf and the merged PDF required by the portal in ../01_documentos/.
"""
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "01_documentos"
TMP = ROOT / "99_temp"
TAG = "Manchester_ssNMR_SAE-032399"

FONT = "Calibri"
ACCENT = RGBColor(0x66, 0x00, 0x99)  # Manchester purple
GREY = RGBColor(0x55, 0x55, 0x55)
NAME = "Dazaet Galicia-Badillo"
CONTACT = "dazaet.galiciab@outlook.com  ·  (+52) 56 5153 4560  ·  Mexico City, Mexico"


# ---------- helpers ----------
def base_doc(margins_cm=1.8):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(margins_cm - 0.2)
        s.left_margin = s.right_margin = Cm(margins_cm)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.08
    return doc


def para(doc, text="", size=None, bold=False, italic=False, color=None,
         align=None, after=None, before=None, indent=None):
    p = doc.add_paragraph()
    if text:
        add_runs(p, text, size=size, bold=bold, italic=italic, color=color)
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    if after is not None:
        pf.space_after = Pt(after)
    if before is not None:
        pf.space_before = Pt(before)
    if indent is not None:
        pf.left_indent = Cm(indent)
    return p


def add_runs(p, text, size=None, bold=False, italic=False, color=None):
    """Text with **bold** markup."""
    for i, part in enumerate(text.split("**")):
        for j, sub in enumerate(part.split("*")):
            if not sub:
                continue
            r = p.add_run(sub)
            r.bold = bold or (i % 2 == 1)
            r.italic = italic or (j % 2 == 1)
            style_run(r, size, color)
    return p


def style_run(r, size, color):
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color


def bottom_border(p, color="660099", sz="6"):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), sz)
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)


def heading(doc, text):
    p = para(doc, text.upper(), size=11, bold=True, color=ACCENT, before=9, after=3)
    bottom_border(p)
    p.paragraph_format.keep_with_next = True
    return p


def dated(doc, left, right, after=1):
    p = doc.add_paragraph()
    sec = doc.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    add_runs(p, left)
    r = p.add_run("\t" + right)
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text, indent=0.45):
    p = doc.add_paragraph(style="List Bullet")
    add_runs(p, text)
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(1.5)
    return p


# ---------- CV ----------
PUBS = [
    "Enriquez-Palacios, E.; **Galicia-Badillo, D.**; Flores-Alamo, M.; Belmonte-Vazquez, J. L. Imidazole-based Photoswitches by Design and their use for Data Encryption. *ACS Appl. Mater. Interfaces* **2026**, 18 (35), 47895-47906.",
    "**Galicia-Badillo, D.**; Hernández-Morales, E. A.; López-Herrera, Y.; Hernández-Santiago, E.; Blahut, J.; Toscano, R. A.; Dračínský, M.; Rodríguez-Molina, B. Dissecting Molecular and Macroscopic Dynamics in a Jumping Solvate. *Mater. Chem. Front.* **2026**, 10 (18), 3305-3313.",
    "**Galicia-Badillo, D.**; Cisneros, J.; Blahut, J.; Acosta-Vera, A.; García-Márquez, A.; Desirena, H.; Bartůnková, E.; García-González, M. C.; Chávez-Serrano, I.; Rodríguez-Hernández, J.; Dračínský, M.; Rodríguez-Molina, B. Integration of a Molecular Rotor and a Photoswitch in a Metal-Organic Framework. *ACS Mater. Lett.* **2026**, 8 (7), 1995-2001.",
    "Hernández-Morales, E. A.; Barreto-Hernández, D.; **Galicia-Badillo, D.**; Toscano, R. A.; García-Aguilera, M. E.; Rodríguez-Molina, B. Crystals as Rockets: Modulation of the Salient Temperature in Cocrystals by Solvent Mixture Composition. *Angew. Chem. Int. Ed.* **2026**, e26010.",
    "Álvarez-Castillo, Á. D.; Juárez-Morales, A.; **Galicia-Badillo, D.**; Maldonado-Domínguez, M.; Heredia-Peñaloza, J. M.; Flores-Álamo, M.; Belmonte-Vázquez, J. L. Rational Modulation of Excited-State Intramolecular Proton Transfer, Aggregation-Induced Emission, and Intramolecular Motion on Polysubstituted Imidazoles. *ChemPhotoChem* **2026**, 10, e70099. (Front cover)",
    "Castañón-García, M.; López-Mendoza, P.; **Galicia-Badillo, D.**; Miranda, L. D. Synthesis of Non-Symmetrical 3,3'-Bicoumarins: Total Synthesis of Arteminorin C, 3,3'-Biisofraxidin and Biscopoletin. *Org. Lett.* **2025**, 27 (21), 5349-5354.",
    "Belmonte-Vázquez, J. L.; Karothu, D. P.; Fleischer, C. H.; **Galicia-Badillo, D.**; Maldonado-Domínguez, M.; Schurko, R. W.; Li, L.; Naumov, P.; Rodríguez-Molina, B. Low-Temperature Dynamics at Nano- and Macroscales: Organic Crystal That Exhibits Low-Temperature Molecular Motion and the Thermosalient Effect. *Chem. Mater.* **2025**, 37, 3373-3383.",
    "**Galicia-Badillo, D.**; Rodríguez-Molina, B. On the Flexibility of 1,4-Bis(phenylethynyl)benzene Butyl Ester Organic Crystal Featuring Mobile Components. *Tetrahedron* **2024**, 167, 134245.",
    "**Galicia-Badillo, D.**; Belmonte-Vázquez, J. L.; Rodríguez, M.; Rodríguez-Molina, B.; García-González, M. C. Aggregation-Induced Emission Enhancement and Solid-State Photoswitching of Crystalline Carbazole N-Salicylidene Anilines. *ACS Omega* **2024**, 9 (36), 38015-38022.",
    "Acosta-Vera, A.; **Galicia-Badillo, D.**; Navarro-Huerta, A.; Stracke, K.; Gómez-Vidales, V.; Rodríguez-Hernández, J.; Evans, J. D.; Rodríguez-Molina, B. Modulation of the Photophysics and Internal Dynamics in a Zr Metal-Organic Framework by the Inclusion of Fluorinated Guests. *ACS Mater. Lett.* **2024**, 6 (9), 4395-4401.",
    "**Galicia-Badillo, D.**; García-Vargas, P. N.; Rodríguez-Molina, B. The Integration of Multiple Components for Multi-Responsive Dynamic Framework Materials (Preview). *Matter* **2024**, 7 (8), 2766-2768.",
    "Hernández-Santiago, E.; Espinosa-Rocha, J.; **Galicia-Badillo, D.**; Rodríguez-Molina, B. Local and Global Dynamics in Organic Soft Porous Crystals (Review). *Matter* **2024**, 7 (2), 408-429.",
    "Rodríguez-Molina, M.; **Galicia-Badillo, D.**; Cetina-Mancilla, E.; Cárdenas, J.; Olvera, L. I.; Toscano, R. A.; Rodríguez-Molina, B.; Zolotukhin, M. G. 9-Trifluoromethylxanthenediols: Synthesis and Supramolecular Motifs. *ACS Omega* **2022**, 7 (16), 13520-13528.",
    "Torres-Huerta, A.; **Galicia-Badillo, D.**; Aguilar-Granda, A.; Bryant, J. T.; Uribe-Romo, F. J.; Rodríguez-Molina, B. Multiple Rotational Rates in a Guest-Loaded, Amphidynamic Zirconia Metal-Organic Framework. *Chem. Sci.* **2020**, 11 (42), 11579-11583.",
]


def add_pub(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.7)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{n}.\t")
    r.font.size = Pt(9.5)
    # handle **bold** and *italic*
    for i, chunk in enumerate(text.split("**")):
        for j, sub in enumerate(chunk.split("*")):
            if not sub:
                continue
            rr = p.add_run(sub)
            rr.bold = i % 2 == 1
            rr.italic = j % 2 == 1
            rr.font.size = Pt(9.5)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.7))


def build_cv():
    doc = base_doc()
    p = para(doc, NAME, size=22, bold=True, color=ACCENT, after=0)
    para(doc, "PhD in Chemical Sciences  ·  Solid-state NMR, X-ray diffraction and dynamics of molecular crystals",
         size=11, color=GREY, after=1)
    para(doc, CONTACT, size=10, after=0)
    para(doc, "Current affiliation: Instituto Mexicano de la Propiedad Industrial (IMPI), Mexico City",
         size=10, color=GREY, after=2)

    heading(doc, "Profile")
    para(doc,
         "Materials chemist trained at the Institute of Chemistry, UNAM, where my PhD used variable-temperature "
         "²H and ¹³C solid-state NMR to measure molecular motion in organic crystals and metal-organic frameworks. "
         "I analyse NMR data by lineshape simulation and relaxation analysis and interpret it together with "
         "variable-temperature single-crystal and powder X-ray diffraction, thermal analysis and DFT. "
         "Author of 14 publications, with collaborators "
         "in the Czech Republic, the United States, China and the United Arab Emirates.", after=2)

    heading(doc, "Research Experience")
    dated(doc, "**Doctoral researcher**, Institute of Chemistry, UNAM (Prof. Braulio Rodríguez-Molina)", "2022 - 2026")
    for b in [
        "Acquired and analysed variable-temperature ²H and ¹³C solid-state NMR data to determine rotational "
        "dynamics of mobile components in organic crystals and Zr-based MOFs, using spectral (lineshape) "
        "simulation and relaxation analysis.",
        "Correlated NMR dynamics with variable-temperature SCXRD/PXRD structures to "
        "separate molecular motion from anisotropic thermal expansion in a thermosalient solvate "
        "(*Mater. Chem. Front.* 2026) and to show how a photoswitch and guest uptake modulate a molecular rotor "
        "inside a MOF (*ACS Mater. Lett.* 2026), in collaboration with the NMR group of Dr Martin Dračínský "
        "(IOCB Prague).",
        "Characterised molecular crystals, cocrystals and solvates by DSC, TGA and hot-stage microscopy, "
        "relating phase transitions, salient behaviour and mechanical flexibility to crystal structure.",
        "Ran periodic DFT (CP2K) and molecular dynamics (LAMMPS) on HPC clusters; wrote two successful "
        "proposals for national supercomputing allocations (LANCAD, 2024).",
        "Synthesised organic molecular rotors and photoswitches (including Schlenk techniques) and MOFs "
        "(solvothermal synthesis and post-synthetic modification); grew single crystals for diffraction.",
    ]:
        bullet(doc, b)
    dated(doc, "**M.Sc. and undergraduate researcher**, Institute of Chemistry, UNAM", "2019 - 2022", after=1)
    bullet(doc, "Studied guest-dependent rotational dynamics in a zirconia MOF (*Chem. Sci.* 2020) and "
                "supramolecular motifs in fluorinated xanthenediols (*ACS Omega* 2022).")

    heading(doc, "Technical Skills")
    for lab, txt in [
        ("Solid-state NMR", "variable-temperature ²H and ¹³C; lineshape simulation; relaxation analysis; "
                            "data processing in MestReNova."),
        ("Diffraction", "SCXRD and PXRD, including variable-temperature studies; structure solution and "
                        "refinement (SHELX); CCDC suite, CrystalDiffract."),
        ("Thermal and optical", "DSC, TGA, hot-stage microscopy; UV-Vis absorption; steady-state and "
                                "time-resolved fluorescence."),
        ("Computation", "DFT in Gaussian16 (geometry optimisation, rotational energy profiles, PES); periodic "
                        "DFT in CP2K; molecular dynamics in LAMMPS; HPC job workflows; AI-assisted data analysis."),
        ("Synthesis", "multistep organic synthesis, air- and moisture-sensitive chemistry, crystallisation of "
                      "molecular crystals, cocrystals and solvates; MOF synthesis and post-synthetic modification."),
        ("Laboratory", "training of students, purchasing and inventory, laboratory organisation and safety."),
    ]:
        p = doc.add_paragraph()
        add_runs(p, f"**{lab}:** {txt}")
        p.paragraph_format.space_after = Pt(2)

    heading(doc, "Education")
    dated(doc, "**PhD in Chemical Sciences** (with honours), UNAM, Institute of Chemistry", "2022 - 2026", after=0)
    para(doc, "Thesis: Molecular and Macroscopic Motion in Organic Crystals and Metal-Organic Frameworks. "
              "Supervisor: Prof. Braulio Rodríguez-Molina.", size=10, color=GREY, after=3)
    dated(doc, "**M.Sc. in Chemical Sciences**, UNAM, Institute of Chemistry", "2020 - 2022", after=3)
    dated(doc, "**B.Sc. in Chemistry**, UNAM, Faculty of Chemistry", "2015 - 2020", after=2)

    heading(doc, "Publications")
    para(doc, "14 publications: 12 research articles, 1 review and 1 preview (2 in *Matter*).",
         size=10, color=GREY, after=3)
    for i, t in enumerate(PUBS):
        add_pub(doc, len(PUBS) - i, t)

    heading(doc, "Honours, Awards and Funding")
    for left, right in [
        ("PhD awarded with honours, UNAM", "2026"),
        ("Nominated for the Alfonso Caso Medal, UNAM (outcome pending)", "2026"),
        ("Best Presentation Award, IX Seminario Regional de Materiales Avanzados, UAEH, Mexico", "2025"),
        ("Two supercomputing allocations, LANCAD (proposals written as lead applicant)", "2024"),
        ("Best Poster Award, DynaMIC23 (Dynamic Materials, Crystals and Phenomena), Fribourg, Switzerland", "2023"),
        ("Doctoral and Master's scholarships, SECIHTI (Mexico)", "2020, 2022"),
        ("Institute of Chemistry Undergraduate Scholarship for best undergraduate project, UNAM", "2019"),
    ]:
        dated(doc, left, right, after=1)

    heading(doc, "Mentoring and Collaboration")
    bullet(doc, "Co-supervised 4 bachelor's theses (2023 - present) on MOF synthesis, crystal dynamics and "
                "luminescent materials; supervised 9 undergraduate research placements (2022 - 2025) and "
                "1 high-school student.")
    bullet(doc, "Contributed experimental work to collaborations with groups in the Czech Republic (IOCB Prague), "
                "the United States, China and the United Arab Emirates.")

    heading(doc, "Conferences and Outreach")
    bullet(doc, "15 posters (6 international) and 9 oral presentations, including 4 invited lectures.")
    bullet(doc, "Outreach: article in Boletín IQ-UNAM (2022), 2 public lectures for high-school audiences "
                "(2024 - 2025), media consultation for UNAM Global (2022).")

    heading(doc, "Referees")
    for t in [
        "**Prof. Braulio Rodríguez-Molina**, PhD supervisor. Institute of Chemistry, UNAM, Mexico. "
        "brodriguez@iquimica.unam.mx · (+52) 55 6706 9475",
        "**Dr Aaron Torres-Huerta**, research collaborator. "
        "aaron.torres.huerta@ulb.be · (+32) 489 34 15 69",
    ]:
        p = doc.add_paragraph()
        add_runs(p, t)
        p.paragraph_format.space_after = Pt(3)

    path = OUT / f"cv-dazaet_galicia_{TAG}.docx"
    doc.save(path)
    return path


# ---------- Covering letter ----------
LETTER = [
    "I am writing to apply for the position of Research Associate in Solid-State NMR Spectroscopy "
    "(ref. SAE-032399) within the EPSRC Programme Grant on Molecular Solid Solutions. I completed my PhD in "
    "Chemical Sciences at the Institute of Chemistry, UNAM (Mexico), in 2026 under the supervision of "
    "Prof. Braulio Rodríguez-Molina, and the degree was awarded with honours.",

    "My doctoral research used solid-state NMR to measure molecular motion in organic crystals and "
    "metal-organic frameworks and to connect that motion with crystal structure and macroscopic behaviour. "
    "I acquired variable-temperature ²H and ¹³C solid-state NMR data and analysed them by lineshape simulation "
    "and relaxation analysis to obtain rotational rates and activation energies. In a pyrrolo[3,2-b]pyrrole "
    "solvate that jumps on heating (Mater. Chem. Front., 2026), I combined ²H NMR with variable-temperature "
    "single-crystal diffraction to separate the molecular dynamics from the anisotropic thermal expansion that "
    "drives the jump. In a zirconium MOF (ACS Mater. Lett., 2026), I used the same combination to follow how a "
    "photoswitch and guest uptake change the motion of a molecular rotor. Both projects were carried out with "
    "Dr Martin Dračínský and Dr Jan Blahut at IOCB Prague, and working with them taught me to use NMR and "
    "diffraction data as checks on each other.",

    "Several of the desirable criteria describe work I do routinely. I characterise molecular crystals, "
    "cocrystals and solvates by SCXRD, PXRD, DSC, TGA and hot-stage microscopy. I run DFT in Gaussian16 and "
    "periodic DFT in CP2K, use LAMMPS for molecular dynamics, and wrote two successful proposals for national "
    "supercomputing time. Multicomponent molecular crystals are already part of my work: I contributed to a study "
    "in which the salient temperature of cocrystals was tuned through the composition of the crystallisation "
    "solvent mixture (Angew. Chem. Int. Ed., 2026). Diffraction describes the average structure of a molecular "
    "solid solution, while solid-state NMR can show how each component is distributed and how it moves locally. "
    "I would very much like to work on that question within MoSS.",

    "My hands-on experience so far is with ²H and ¹³C. I am keen to extend it to ¹H, ¹⁵N, ¹⁹F and ¹⁷O "
    "experiments and to DNP-enhanced MAS, and the opportunity to learn these methods in your group is one of "
    "the main reasons for this application. During my PhD I learned periodic DFT and molecular dynamics "
    "from scratch and now use both in my projects, and I would approach new NMR methods in the same way.",

    "I enjoy working in teams spread across institutions. My 14 publications include collaborations with groups "
    "in the Czech Republic, the United States, China and the United Arab Emirates. I co-supervised four bachelor's "
    "theses and nine undergraduate research placements, and I received the Best Poster Award at DynaMIC23 in "
    "Fribourg. A consortium shared between Durham, Leeds and Manchester, with placements between sites, suits "
    "the way I like to work.",

    "Thank you for considering my application. I would be glad to discuss it at interview.",
]


def build_letter():
    doc = base_doc(margins_cm=2.0)
    para(doc, NAME, size=18, bold=True, color=ACCENT, after=0)
    para(doc, CONTACT, size=10, color=GREY, after=0)
    bottom_border(doc.paragraphs[-1])
    para(doc, "", after=4)
    para(doc, "10 October 2026", after=8)
    for line in ["Dr Daniel Lee", "Department of Chemical Engineering, School of Engineering",
                 "The University of Manchester", "Manchester M13 9PL, United Kingdom"]:
        para(doc, line, after=0)
    para(doc, "", after=4)
    para(doc, "**Re: Research Associate in Solid-State NMR Spectroscopy, ref. SAE-032399**", after=8)
    para(doc, "Dear Dr Lee,", after=6)
    for t in LETTER:
        p = para(doc, t, after=5)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para(doc, "Yours sincerely,", after=12)
    para(doc, "Dazaet Galicia-Badillo, PhD", bold=True, after=0)
    path = OUT / f"coverletter-dazaet_galicia_{TAG}.docx"
    doc.save(path)
    return path


def to_pdf(docx_path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(docx_path)],
                   check=True, capture_output=True)
    for junk in OUT.glob("*.tmp"):
        shutil.move(str(junk), TMP / junk.name)
    return docx_path.with_suffix(".pdf")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    TMP.mkdir(exist_ok=True)
    letter_pdf = to_pdf(build_letter())
    cv_pdf = to_pdf(build_cv())
    merged = OUT / f"Galicia-Badillo_Dazaet_SAE-032399_CoverLetter_CV.pdf"
    w = PdfWriter()
    for pdf in (letter_pdf, cv_pdf):
        w.append(str(pdf))
    w.add_metadata({"/Title": "Covering letter and CV, Dazaet Galicia-Badillo, SAE-032399",
                    "/Author": NAME})
    w.write(str(merged))
    print(merged, merged.stat().st_size // 1024, "KB")
